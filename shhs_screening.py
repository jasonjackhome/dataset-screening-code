import os
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

# ============================================================
# 1. File paths
# ============================================================

shhs1_file_path = r"F:\shhs\datasets\shhs1-dataset-0.20.0.csv"
shhs2_file_path = r"F:\shhs\datasets\shhs2-dataset-0.20.0.csv"

output_dir = r"D:\Users\shhs\shhs_screening"

os.makedirs(output_dir, exist_ok=True)

output_s1_xlsx = os.path.join(
    output_dir,
    "SHHS1_exclusion_list.xlsx"
)

output_s2_xlsx = os.path.join(
    output_dir,
    "SHHS2_exclusion_list.xlsx"
)


# ============================================================
# 2. Helper functions
# ============================================================

def normalize_nsrrid(series):
    """
    Normalize NSRR subject IDs for matching between SHHS1 and SHHS2.
    The original nsrrid values are not modified.
    """
    return (
        series.astype("string")
        .str.strip()
        .str.replace(r"\.0$", "", regex=True)
    )


def to_number(value):
    """
    Safely convert a value to numeric.
    """
    try:
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def reason_ahi(value):
    """
    AHI must be < 5 events/h to satisfy the screening criterion.
    """
    x = to_number(value)

    if pd.isna(x):
        return "Missing AHI information (ahi_a0h3a)"

    if x >= 5:
        return f"AHI >= 5 events/h (ahi_a0h3a={x:g})"

    return None


def reason_zero_required(
        value,
        variable_name,
        positive_reason,
        code8_reason=None):
    """
    Evaluate variables that must be explicitly coded as 0.
    """

    x = to_number(value)

    if pd.isna(x):
        return f"Missing information ({variable_name})"

    if x == 0:
        return None

    if x == 1:
        return positive_reason

    if x == 8:
        if code8_reason is not None:
            return code8_reason
        else:
            return f"Unknown/not applicable status ({variable_name}=8)"

    return (
        f"Status not confirmed as negative "
        f"({variable_name}={value})"
    )


def join_reasons(reasons):
    """
    Combine multiple exclusion reasons for one participant.
    """
    reasons = [x for x in reasons if x is not None]

    if len(reasons) == 0:
        return ""

    return "; ".join(reasons)


# ============================================================
# 3. Read datasets
# ============================================================

print("=" * 80)
print("Reading SHHS data...")
print("=" * 80)

df1 = pd.read_csv(
    shhs1_file_path,
    low_memory=False,
    encoding="latin1"
)

df2 = pd.read_csv(
    shhs2_file_path,
    low_memory=False,
    encoding="latin1"
)

print(f"SHHS1 records: {len(df1)}")
print(f"SHHS2 records: {len(df2)}")

# ============================================================
# 4. Check required variables
# ============================================================

required_s1 = [
    "nsrrid",
    "shhs1_psg",
    "ahi_a0h3a",
    "abnoreeg",
    "stroke15",
    "prev_hx_stroke",
    "afib",
    "hf15"
]

required_s2 = [
    "nsrrid",
    "shhs2_psg",
    "ahi_a0h3a",
    "abnoreeg",
    "afib",
    "alzh2",
    "comm"
]

missing_s1 = [x for x in required_s1 if x not in df1.columns]
missing_s2 = [x for x in required_s2 if x not in df2.columns]

if missing_s1:
    raise KeyError(
        f"Missing required SHHS1 variables: {missing_s1}"
    )

if missing_s2:
    raise KeyError(
        f"Missing required SHHS2 variables: {missing_s2}"
    )

# ============================================================
# 5. SHHS1 screening
# ============================================================

print("\n" + "=" * 80)
print("Processing SHHS1...")
print("=" * 80)

s1 = df1.copy()

s1["nsrrid_key"] = normalize_nsrrid(s1["nsrrid"])


# ------------------------------------------------------------
# Generate exclusion reasons for SHHS1 participants
# ------------------------------------------------------------

def get_s1_exclusion_reason(row):
    reasons = []

    # PSG availability
    shhs1_psg = to_number(row["shhs1_psg"])

    if pd.isna(shhs1_psg) or shhs1_psg != 1:
        reasons.append(
            "PSG availability criterion not met (shhs1_psg)"
        )

    # AHI
    reasons.append(
        reason_ahi(row["ahi_a0h3a"])
    )

    # Abnormal EEG
    reasons.append(
        reason_zero_required(
            row["abnoreeg"],
            "abnoreeg",
            "Abnormal EEG observed during PSG",
            "Abnormal EEG status not applicable/unknown "
            "(abnoreeg=8)"
        )
    )

    # MD reported stroke
    reasons.append(
        reason_zero_required(
            row["stroke15"],
            "stroke15",
            "MD-reported stroke",
            "MD-reported stroke status unknown "
            "(stroke15=8)"
        )
    )

    # Previous history of stroke
    reasons.append(
        reason_zero_required(
            row["prev_hx_stroke"],
            "prev_hx_stroke",
            "Previous history of stroke"
        )
    )

    # AF / flutter
    reasons.append(
        reason_zero_required(
            row["afib"],
            "afib",
            "Atrial fibrillation or flutter"
        )
    )

    # Heart failure
    reasons.append(
        reason_zero_required(
            row["hf15"],
            "hf15",
            "MD-reported heart failure",
            "Heart failure status unknown "
            "(hf15=8)"
        )
    )

    return join_reasons(reasons)


s1["Exclusion_reason"] = s1.apply(
    get_s1_exclusion_reason,
    axis=1
)

s1["Final_status"] = np.where(
    s1["Exclusion_reason"] == "",
    "Included",
    "Excluded"
)

# ============================================================
# 6. Identify participants retained after SHHS1 screening
# ============================================================

healthy_s1 = s1[
    s1["Final_status"] == "Included"
    ].copy()

print(
    f"SHHS1 included subjects: "
    f"{len(healthy_s1)} / {len(s1)}"
)

print(
    f"SHHS1 excluded subjects: "
    f"{len(s1) - len(healthy_s1)}"
)

# ============================================================
# 7. SHHS2 screening with SHHS1 baseline status mapping
# ============================================================

print("\n" + "=" * 80)
print("Processing SHHS2...")
print("=" * 80)

s2 = df2.copy()

s2["nsrrid_key"] = normalize_nsrrid(
    s2["nsrrid"]
)

# ------------------------------------------------------------
# Map SHHS1 screening results to SHHS2
# ------------------------------------------------------------

s1_status_map = dict(
    zip(
        s1["nsrrid_key"],
        s1["Final_status"]
    )
)

s1_reason_map = dict(
    zip(
        s1["nsrrid_key"],
        s1["Exclusion_reason"]
    )
)

s2["SHHS1_status"] = (
    s2["nsrrid_key"]
    .map(s1_status_map)
)

s2["SHHS1_exclusion_reason"] = (
    s2["nsrrid_key"]
    .map(s1_reason_map)
)


# ------------------------------------------------------------
# Generate SHHS2-specific exclusion reasons
# ------------------------------------------------------------

def reason_s2_recording_qc(value):
    """
    In addition, the suitability of the final candidate EEG recordings
    was further reviewed because the present study relies on frequency-band-specific
    EEG representations. For one otherwise eligible SHHS2 recording (nsrrid: 200908),
    the 'comm' annotation indicated that the EEG signal was considered acceptable
    only after application of a 15-Hz low-pass filter. Because our experiments use
    EEG signals filtered within 0.3–35 Hz and explicitly include a beta-band input
    of 13–35 Hz, this additional 15-Hz low-pass filtering would alter the spectral
    information required for the predefined frequency-band analysis.
    Following review by experienced sleep experts, this recording was therefore
    excluded from the final analytic dataset.
    """
    if pd.isna(value):
        return None

    text = str(value).lower()
    compact = text.replace(" ", "")

    if ("eeg" in text) and ("15hzlowpass" in compact):
        return (
            "Recording-level EEG spectral incompatibility: "
            "SHHS quality-control note documented a 15-Hz "
            "low-pass filter applied to EEG"
        )

    return None


def get_s2_reason(row):
    reasons = []

    # AHI
    reasons.append(
        reason_ahi(
            row["ahi_a0h3a"]
        )
    )

    # Abnormal EEG
    reasons.append(
        reason_zero_required(
            row["abnoreeg"],
            "abnoreeg",
            "Abnormal EEG observed during PSG",
            "Abnormal EEG status not applicable/unknown "
            "(abnoreeg=8)"
        )
    )

    # AF / flutter
    reasons.append(
        reason_zero_required(
            row["afib"],
            "afib",
            "Atrial fibrillation or flutter"
        )
    )

    # Alzheimer-related medication
    reasons.append(
        reason_zero_required(
            row["alzh2"],
            "alzh2",
            "Use of acetylcholinesterase inhibitors "
            "for Alzheimer's disease within two weeks "
            "of the SHHS2 visit",
            "Alzheimer-related medication status "
            "unknown/not applicable (alzh2=8)"
        )
    )

    # Recording-level EEG quality control
    reasons.append(
        reason_s2_recording_qc(
            row["comm"]
        )
    )

    return join_reasons(reasons)


s2["SHHS2_exclusion_reason"] = s2.apply(
    get_s2_reason,
    axis=1
)


# ============================================================
# 8. Determine final SHHS2 screening status
# ============================================================

def determine_s2_final_status(row):
    # 1. No matching SHHS1 baseline record
    if pd.isna(row["SHHS1_status"]):
        return pd.Series([
            "Excluded",
            "SHHS1 baseline",
            "No matching SHHS1 baseline record"
        ])

    # 2. Participant excluded during SHHS1 screening
    if row["SHHS1_status"] != "Included":

        reason = row["SHHS1_exclusion_reason"]

        if pd.isna(reason) or reason == "":
            reason = "Failed SHHS1 baseline screening"

        return pd.Series([
            "Excluded",
            "SHHS1 baseline",
            "Failed SHHS1 baseline criteria: " + reason
        ])

    # 3. SHHS1 criteria satisfied, but no SHHS2 PSG is available
    shhs2_psg = pd.to_numeric(
        pd.Series([row["shhs2_psg"]]),
        errors="coerce"
    ).iloc[0]

    if pd.isna(shhs2_psg) or shhs2_psg != 1:
        return pd.Series([
            "Excluded",
            "SHHS2 PSG availability",
            "No available SHHS2 PSG recording"
        ])

    # 4. SHHS2 PSG available, but clinical/QC criteria are not satisfied
    if row["SHHS2_exclusion_reason"] != "":
        return pd.Series([
            "Excluded",
            "SHHS2 follow-up",
            row["SHHS2_exclusion_reason"]
        ])

    # 5. All screening criteria satisfied
    return pd.Series([
        "Included",
        "None",
        ""
    ])


s2[
    [
        "Final_status",
        "Exclusion_stage",
        "Exclusion_reason"
    ]
] = s2.apply(
    determine_s2_final_status,
    axis=1
)


# ============================================================
# 9. Generate exclusion reasons
# ============================================================

def for_s1_reason(row):
    """
    SHHS1 exclusion reason
    """

    failed_variables = []

    # PSG recording must be available
    shhs1_psg = pd.to_numeric(
        pd.Series([row["shhs1_psg"]]),
        errors="coerce"
    ).iloc[0]

    if pd.isna(shhs1_psg) or shhs1_psg != 1:
        failed_variables.append("'shhs1_psg'")

    # AHI must be < 5 events/h
    ahi = pd.to_numeric(
        pd.Series([row["ahi_a0h3a"]]),
        errors="coerce"
    ).iloc[0]

    if pd.isna(ahi) or ahi >= 5:
        failed_variables.append("'ahi_a0h3a'")

    # All categorical screening variables must be explicitly coded as 0
    categorical_variables = [
        "abnoreeg",
        "stroke15",
        "prev_hx_stroke",
        "afib",
        "hf15"
    ]

    for var in categorical_variables:

        value = pd.to_numeric(
            pd.Series([row[var]]),
            errors="coerce"
        ).iloc[0]

        if pd.isna(value) or value != 0:
            failed_variables.append(f"'{var}'")

    if failed_variables:
        return (
                "Screening criterion not met ("
                + ", ".join(failed_variables)
                + ")"
        )

    return ""


def for_s2_reason(row):
    """
    SHHS2 exclusion reason.
    """

    # --------------------------------------------------------
    # 1. Participant not retained after SHHS1 baseline screening
    # --------------------------------------------------------
    if pd.isna(row["SHHS1_status"]):
        return "Participant not retained after SHHS1 screening"

    if row["SHHS1_status"] != "Included":
        return "Participant not retained after SHHS1 screening"

    # --------------------------------------------------------
    # 2. SHHS2 PSG recording unavailable
    # --------------------------------------------------------
    shhs2_psg = pd.to_numeric(
        pd.Series([row["shhs2_psg"]]),
        errors="coerce"
    ).iloc[0]

    if pd.isna(shhs2_psg) or shhs2_psg != 1:
        return (
            "PSG availability criterion not met "
            "('shhs2_psg')"
        )

    # --------------------------------------------------------
    # 3. SHHS2 clinical screening
    # --------------------------------------------------------
    reasons = []
    failed_variables = []

    ahi = pd.to_numeric(
        pd.Series([row["ahi_a0h3a"]]),
        errors="coerce"
    ).iloc[0]

    if pd.isna(ahi) or ahi >= 5:
        failed_variables.append("'ahi_a0h3a'")

    categorical_variables = [
        "abnoreeg",
        "afib",
        "alzh2"
    ]

    for var in categorical_variables:

        value = pd.to_numeric(
            pd.Series([row[var]]),
            errors="coerce"
        ).iloc[0]

        if pd.isna(value) or value != 0:
            failed_variables.append(f"'{var}'")

    if failed_variables:
        reasons.append(
            "Screening criterion not met ("
            + ", ".join(failed_variables)
            + ")"
        )

    # --------------------------------------------------------
    # 4. Recording-level expert review
    # --------------------------------------------------------
    qc_reason = reason_s2_recording_qc(
        row["comm"]
    )

    if qc_reason is not None:
        reasons.append(
            "Recording unsuitable after expert review"
        )

    return "; ".join(reasons)


# ============================================================
# SHHS1: all participants excluded during baseline screening
# ============================================================

for_s1_excluded = s1[
    s1["Final_status"] == "Excluded"
    ][
    ["nsrrid"]
].copy()

for_s1_excluded[
    "Specific exclusion reasons"
] = s1[
    s1["Final_status"] == "Excluded"
    ].apply(
    for_s1_reason,
    axis=1
).values

for_s1_excluded = for_s1_excluded.rename(
    columns={
        "nsrrid": "Subject ID"
    }
)

# ============================================================
# SHHS2: all records not retained in the final cohort
# The reported reason reflects the stage at which exclusion occurred.
# ============================================================

for_s2_source = s2[
    s2["Final_status"] == "Excluded"
    ].copy()

for_s2_excluded = for_s2_source[
    ["nsrrid"]
].copy()

for_s2_excluded[
    "Specific exclusion reasons"
] = for_s2_source.apply(
    for_s2_reason,
    axis=1
).values

for_s2_excluded = for_s2_excluded.rename(
    columns={
        "nsrrid": "Subject ID"
    }
)

# ============================================================
# 10. Sort exclusion lists by subject ID
# ============================================================

for_s1_excluded = for_s1_excluded.sort_values(
    by="Subject ID"
).reset_index(drop=True)

for_s2_excluded = for_s2_excluded.sort_values(
    by="Subject ID"
).reset_index(drop=True)

# ============================================================
# 11. Save Excel files
# ============================================================

for_s1_excluded.to_excel(
    output_s1_xlsx,
    index=False,
    sheet_name="SHHS1 Exclusion List"
)

for_s2_excluded.to_excel(
    output_s2_xlsx,
    index=False,
    sheet_name="SHHS2 Exclusion List"
)


# ============================================================
# 12. Format Excel files
# ============================================================


def format_excel(file_path):
    wb = load_workbook(file_path)
    ws = wb.active

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = ws.dimensions

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 55

    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(
                name="Times New Roman",
                size=11,
                bold=(cell.row == 1)
            )

    for cell in ws[1]:
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(
            horizontal="center",
            vertical="top"
        )

        row[1].alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True
        )

    wb.save(file_path)


format_excel(
    output_s1_xlsx
)

format_excel(
    output_s2_xlsx
)

# ============================================================
# 13. Print summary
# ============================================================

print("\n" + "=" * 80)
print("Screening completed.")
print("=" * 80)

print(
    f"\nSHHS1 exclusion list: "
    f"{len(for_s1_excluded)} subjects"
)

print(
    f"SHHS2 exclusion list: "
    f"{len(for_s2_excluded)} records"
)

print("\nFiles saved:")
print(output_s1_xlsx)
print(output_s2_xlsx)
